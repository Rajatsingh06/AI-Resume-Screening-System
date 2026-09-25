import io
import os
import re
from datetime import date
from pathlib import Path

from flask import Flask, render_template_string, request
from werkzeug.utils import secure_filename

from docx import Document
from pypdf import PdfReader
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity


app = Flask(__name__)

# Maximum total request size: 50 MB
app.config["MAX_CONTENT_LENGTH"] = 50 * 1024 * 1024


ALLOWED_EXTENSIONS = {"pdf", "docx", "txt"}


# -------------------------------------------------------------------
# Skill dictionary
# -------------------------------------------------------------------

SKILL_ALIASES = {
    "Python": ["python"],
    "Java": ["java"],
    "JavaScript": ["javascript", "js"],
    "TypeScript": ["typescript", "ts"],
    "C++": ["c++", "cpp"],
    "C#": ["c#", "csharp"],
    "Go": ["golang", "go language"],
    "Rust": ["rust"],
    "SQL": ["sql"],
    "NoSQL": ["nosql"],
    "HTML": ["html"],
    "CSS": ["css"],
    "React": ["react", "react.js", "reactjs"],
    "Angular": ["angular"],
    "Vue": ["vue", "vue.js", "vuejs"],
    "Node.js": ["node.js", "nodejs", "node js"],
    "Express.js": ["express.js", "expressjs"],
    "Django": ["django"],
    "Flask": ["flask"],
    "Spring Boot": ["spring boot"],
    ".NET": [".net", "dotnet", "asp.net"],
    "REST API": ["rest api", "restful api", "rest services"],
    "GraphQL": ["graphql"],
    "Git": ["git"],
    "GitHub": ["github"],
    "Docker": ["docker"],
    "Kubernetes": ["kubernetes", "k8s"],
    "AWS": ["aws", "amazon web services"],
    "Azure": ["azure", "microsoft azure"],
    "Google Cloud": ["google cloud", "gcp"],
    "Linux": ["linux"],
    "Terraform": ["terraform"],
    "Jenkins": ["jenkins"],
    "GitHub Actions": ["github actions"],
    "PostgreSQL": ["postgresql", "postgres"],
    "MySQL": ["mysql"],
    "MongoDB": ["mongodb", "mongo db"],
    "Redis": ["redis"],
    "Pandas": ["pandas"],
    "NumPy": ["numpy"],
    "Scikit-learn": ["scikit-learn", "sklearn"],
    "TensorFlow": ["tensorflow"],
    "PyTorch": ["pytorch"],
    "Machine Learning": ["machine learning", "ml"],
    "Deep Learning": ["deep learning"],
    "NLP": ["nlp", "natural language processing"],
    "Computer Vision": ["computer vision"],
    "Data Science": ["data science"],
    "Data Analysis": ["data analysis", "data analytics"],
    "Power BI": ["power bi"],
    "Tableau": ["tableau"],
    "Excel": ["excel", "microsoft excel"],
    "Apache Spark": ["apache spark", "spark"],
    "Hadoop": ["hadoop"],
    "Airflow": ["airflow", "apache airflow"],
    "Agile": ["agile"],
    "Scrum": ["scrum"],
    "Project Management": ["project management"],
    "Communication": ["communication skills", "verbal communication"],
    "Leadership": ["leadership"],
    "Problem Solving": ["problem solving", "problem-solving"],
}


EDUCATION_PATTERNS = [
    (
        5,
        [
            r"\bph\.?\s*d\b",
            r"\bdoctorate\b",
            r"\bdoctoral\b",
        ],
    ),
    (
        4,
        [
            r"\bmaster'?s\b",
            r"\bmasters\b",
            r"\bm\.?s\.?\b",
            r"\bmba\b",
            r"\bmaster of science\b",
            r"\bmaster of arts\b",
        ],
    ),
    (
        3,
        [
            r"\bbachelor'?s\b",
            r"\bbachelors\b",
            r"\bb\.?s\.?\b",
            r"\bb\.?a\.?\b",
            r"\bbsc\b",
            r"\bba degree\b",
            r"\bbachelor of science\b",
            r"\bbachelor of arts\b",
        ],
    ),
    (
        2,
        [
            r"\bassociate'?s\b",
            r"\bassociates\b",
            r"\ba\.?s\.?\b",
            r"\ba\.?a\.?\b",
        ],
    ),
    (
        1,
        [
            r"\bdiploma\b",
            r"\bcertificate\b",
            r"\bcertification\b",
        ],
    ),
]


# -------------------------------------------------------------------
# Utility functions
# -------------------------------------------------------------------

def get_extension(filename):
    return Path(filename).suffix.lower().replace(".", "")


def is_allowed_file(filename):
    return get_extension(filename) in ALLOWED_EXTENSIONS


def normalize_text(text):
    if not text:
        return ""
    return re.sub(r"\s+", " ", text.lower()).strip()


def extract_text_from_file(file_storage):
    """
    Extract text from PDF, DOCX, or TXT upload.
    """
    filename = file_storage.filename or ""
    extension = get_extension(filename)
    raw_data = file_storage.read()

    if extension == "txt":
        return raw_data.decode("utf-8", errors="ignore")

    if extension == "pdf":
        reader = PdfReader(io.BytesIO(raw_data))
        pages = []

        for page in reader.pages:
            try:
                pages.append(page.extract_text() or "")
            except Exception:
                pages.append("")

        return "\n".join(pages)

    if extension == "docx":
        document = Document(io.BytesIO(raw_data))
        paragraphs = [paragraph.text for paragraph in document.paragraphs]

        # Also collect text from tables
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    paragraphs.append(cell.text)

        return "\n".join(paragraphs)

    raise ValueError("Unsupported file type")


def alias_exists(text, alias):
    """
    Detect a skill alias in a text string.
    """
    alias = alias.lower().strip()
    escaped_alias = re.escape(alias)

    # Allow flexible spaces in multi-word aliases
    escaped_alias = escaped_alias.replace(r"\ ", r"\s+")

    # Use word boundaries for simple words such as Python or SQL.
    # Avoid word boundaries around symbols such as C++, .NET, and C#.
    if re.fullmatch(r"[a-z0-9]+", alias):
        pattern = rf"\b{escaped_alias}\b"
    else:
        pattern = escaped_alias

    return re.search(pattern, text, flags=re.IGNORECASE) is not None


def extract_skills(text):
    """
    Extract known skills from text using an alias dictionary.
    """
    normalized = normalize_text(text)
    found_skills = []

    for canonical_name, aliases in SKILL_ALIASES.items():
        if any(alias_exists(normalized, alias) for alias in aliases):
            found_skills.append(canonical_name)

    return sorted(found_skills)


def extract_education_level(text):
    """
    Returns an education level from 0 to 5.
    """
    normalized = normalize_text(text)
    highest_level = 0

    for level, patterns in EDUCATION_PATTERNS:
        for pattern in patterns:
            if re.search(pattern, normalized, flags=re.IGNORECASE):
                highest_level = max(highest_level, level)

    return highest_level


def education_label(level):
    labels = {
        0: "Not detected",
        1: "Certificate/Diploma",
        2: "Associate Degree",
        3: "Bachelor's Degree",
        4: "Master's Degree",
        5: "Doctorate",
    }
    return labels.get(level, "Not detected")


def extract_years_of_experience(text):
    """
    Extract approximate years of experience.

    Examples detected:
        3 years experience
        5+ years
        2018 - 2022
        2020 to Present
    """
    normalized = normalize_text(text)
    values = []

    # Detect ranges such as 3-5 years or 3 to 5 years
    range_pattern = (
        r"(\d+(?:\.\d+)?)\s*(?:-|to)\s*"
        r"(\d+(?:\.\d+)?)\s*(?:years?|yrs?)"
    )

    for match in re.finditer(range_pattern, normalized):
        first_value = float(match.group(1))
        second_value = float(match.group(2))
        values.append(max(first_value, second_value))

    # Detect single values such as 5 years or 5+ years
    single_pattern = r"(\d+(?:\.\d+)?)\s*\+?\s*(?:years?|yrs?)"

    for match in re.finditer(single_pattern, normalized):
        values.append(float(match.group(1)))

    explicit_years = max(values) if values else 0.0

    # Estimate experience from date ranges
    date_pattern = (
        r"((?:19|20)\d{2})\s*"
        r"(?:-|–|—|to)\s*"
        r"((?:19|20)\d{2}|present|current)"
    )

    date_durations = []

    for match in re.finditer(date_pattern, normalized):
        start_year = int(match.group(1))
        end_value = match.group(2)

        if end_value in {"present", "current"}:
            end_year = date.today().year
        else:
            end_year = int(end_value)

        if end_year >= start_year:
            date_durations.append(end_year - start_year)

    estimated_from_dates = sum(date_durations)

    return round(max(explicit_years, estimated_from_dates), 1)


def extract_required_years(job_description):
    """
    Extract the highest stated experience requirement from the job description.
    """
    return extract_years_of_experience(job_description)


def get_candidate_name(filename, text):
    """
    Attempt to identify a candidate name from the first lines.
    Otherwise use the uploaded filename.
    """
    lines = [line.strip() for line in text.splitlines() if line.strip()]

    for line in lines[:10]:
        if len(line) > 80:
            continue

        if "@" in line:
            continue

        if re.search(r"\d{3,}", line):
            continue

        if len(line.split()) <= 5:
            excluded_headers = {
                "resume",
                "curriculum vitae",
                "professional summary",
                "work experience",
                "education",
                "skills",
                "profile",
            }

            if line.lower() not in excluded_headers:
                return line

    safe_name = secure_filename(filename)
    stem = Path(safe_name).stem
    stem = stem.replace("_", " ").replace("-", " ").strip()

    return stem.title() if stem else "Unnamed Candidate"


def calculate_experience_score(candidate_years, required_years):
    if required_years <= 0:
        return 1.0

    return min(candidate_years / required_years, 1.0)


def calculate_education_score(candidate_level, required_level):
    if required_level <= 0:
        return 1.0

    if candidate_level >= required_level:
        return 1.0

    if candidate_level > 0:
        return 0.5

    return 0.0


def analyze_candidates(job_description, candidates):
    """
    Analyze and rank candidates using:

    - Skill matching
    - TF-IDF cosine similarity
    - Experience matching
    - Education matching
    """
    normalized_job_description = normalize_text(job_description)

    required_skills = extract_skills(job_description)
    required_years = extract_required_years(job_description)
    required_education = extract_education_level(job_description)

    candidate_texts = [
        normalize_text(candidate["text"])
        for candidate in candidates
    ]

    # TF-IDF semantic similarity
    all_documents = [normalized_job_description] + candidate_texts

    try:
        vectorizer = TfidfVectorizer(
            stop_words="english",
            ngram_range=(1, 2),
            max_features=6000,
        )

        matrix = vectorizer.fit_transform(all_documents)

        if len(candidates) > 0:
            similarities = cosine_similarity(
                matrix[0:1],
                matrix[1:],
            )[0]
        else:
            similarities = []

    except ValueError:
        similarities = [0.0 for _ in candidates]

    analyzed = []

    for index, candidate in enumerate(candidates):
        candidate_text = candidate["text"]

        candidate_skills = extract_skills(candidate_text)
        matched_skills = sorted(
            set(required_skills).intersection(candidate_skills)
        )
        missing_skills = sorted(
            set(required_skills).difference(candidate_skills)
        )

        if required_skills:
            skill_score = (
                len(matched_skills) / len(required_skills)
            )
        else:
            skill_score = 0.0

        semantic_score = (
            float(similarities[index]) if index < len(similarities) else 0.0
        )

        candidate_years = extract_years_of_experience(candidate_text)
        experience_score = calculate_experience_score(
            candidate_years,
            required_years,
        )

        candidate_education = extract_education_level(candidate_text)
        education_score = calculate_education_score(
            candidate_education,
            required_education,
        )

        # If the job description contains recognizable skills,
        # use skills as the strongest signal.
        if required_skills:
            final_score = (
                skill_score * 0.55
                + semantic_score * 0.25
                + experience_score * 0.10
                + education_score * 0.10
            )
        else:
            # If no known skills are detected, rely more heavily
            # on semantic similarity.
            final_score = (
                semantic_score * 0.70
                + experience_score * 0.20
                + education_score * 0.10
            )

        result = {
            "name": candidate["name"],
            "filename": candidate["filename"],
            "matched_skills": matched_skills,
            "missing_skills": missing_skills,
            "candidate_skills": candidate_skills,
            "candidate_years": candidate_years,
            "required_years": required_years,
            "candidate_education": education_label(candidate_education),
            "required_education": education_label(required_education),
            "skill_score": round(skill_score * 100, 1),
            "semantic_score": round(semantic_score * 100, 1),
            "experience_score": round(experience_score * 100, 1),
            "education_score": round(education_score * 100, 1),
            "final_score": round(final_score * 100, 1),
        }

        analyzed.append(result)

    analyzed.sort(
        key=lambda candidate: candidate["final_score"],
        reverse=True,
    )

    for rank, candidate in enumerate(analyzed, start=1):
        candidate["rank"] = rank

    return {
        "results": analyzed,
        "required_skills": required_skills,
        "required_years": required_years,
        "required_education": education_label(required_education),
    }


# -------------------------------------------------------------------
# HTML template
# -------------------------------------------------------------------

HTML_TEMPLATE = """
<!DOCTYPE html>
<html lang="en">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>AI Resume Screening System</title>

    <style>
        :root {
            --primary: #2563eb;
            --primary-dark: #1d4ed8;
            --secondary: #0f172a;
            --background: #f8fafc;
            --card: #ffffff;
            --border: #e2e8f0;
            --muted: #64748b;
            --success: #16a34a;
            --warning: #d97706;
            --danger: #dc2626;
        }

        * {
            box-sizing: border-box;
        }

        body {
            margin: 0;
            font-family: Arial, Helvetica, sans-serif;
            background: var(--background);
            color: var(--secondary);
        }

        header {
            background: linear-gradient(
                135deg,
                #1d4ed8,
                #2563eb,
                #38bdf8
            );
            color: white;
            padding: 42px 20px;
        }

        .header-content {
            max-width: 1180px;
            margin: 0 auto;
        }

        header h1 {
            margin: 0 0 12px;
            font-size: 34px;
        }

        header p {
            max-width: 760px;
            margin: 0;
            line-height: 1.6;
            opacity: 0.95;
        }

        .container {
            width: min(1180px, calc(100% - 32px));
            margin: 28px auto 60px;
        }

        .card {
            background: var(--card);
            border: 1px solid var(--border);
            border-radius: 16px;
            padding: 24px;
            box-shadow: 0 8px 25px rgba(15, 23, 42, 0.05);
            margin-bottom: 22px;
        }

        .card h2 {
            margin-top: 0;
            margin-bottom: 8px;
        }

        .muted {
            color: var(--muted);
        }

        .grid {
            display: grid;
            grid-template-columns: 1fr 1fr;
            gap: 20px;
        }

        @media (max-width: 800px) {
            .grid {
                grid-template-columns: 1fr;
            }
        }

        label {
            display: block;
            font-weight: bold;
            margin-bottom: 8px;
        }

        textarea,
        input[type="file"] {
            width: 100%;
            border: 1px solid var(--border);
            border-radius: 10px;
            padding: 12px;
            font-size: 14px;
            background: white;
        }

        textarea {
            min-height: 230px;
            resize: vertical;
            line-height: 1.5;
        }

        input[type="file"] {
            padding: 10px;
        }

        .field {
            margin-bottom: 18px;
        }

        .hint {
            font-size: 13px;
            color: var(--muted);
            margin-top: 7px;
        }

        .button-row {
            display: flex;
            align-items: center;
            gap: 12px;
            flex-wrap: wrap;
        }

        button {
            background: var(--primary);
            color: white;
            border: none;
            border-radius: 9px;
            padding: 13px 22px;
            font-weight: bold;
            cursor: pointer;
            font-size: 14px;
        }

        button:hover {
            background: var(--primary-dark);
        }

        .secondary-button {
            background: #e2e8f0;
            color: #0f172a;
        }

        .secondary-button:hover {
            background: #cbd5e1;
        }

        .alert {
            border-radius: 10px;
            padding: 14px 16px;
            margin-bottom: 18px;
        }

        .alert-error {
            background: #fef2f2;
            border: 1px solid #fecaca;
            color: #991b1b;
        }

        .alert-info {
            background: #eff6ff;
            border: 1px solid #bfdbfe;
            color: #1e40af;
        }

        .summary-grid {
            display: grid;
            grid-template-columns: repeat(4, 1fr);
            gap: 14px;
            margin-top: 18px;
        }

        @media (max-width: 800px) {
            .summary-grid {
                grid-template-columns: repeat(2, 1fr);
            }
        }

        @media (max-width: 450px) {
            .summary-grid {
                grid-template-columns: 1fr;
            }
        }

        .summary-box {
            background: #f8fafc;
            border: 1px solid var(--border);
            border-radius: 12px;
            padding: 16px;
        }

        .summary-box strong {
            display: block;
            font-size: 23px;
            color: var(--primary);
            margin-bottom: 5px;
        }

        .summary-box span {
            color: var(--muted);
            font-size: 13px;
        }

        .candidate-card {
            background: white;
            border: 1px solid var(--border);
            border-radius: 16px;
            padding: 20px;
            margin-bottom: 16px;
        }

        .candidate-top {
            display: flex;
            align-items: center;
            gap: 16px;
            justify-content: space-between;
            flex-wrap: wrap;
        }

        .candidate-info {
            display: flex;
            align-items: center;
            gap: 14px;
        }

        .rank {
            width: 45px;
            height: 45px;
            display: flex;
            align-items: center;
            justify-content: center;
            border-radius: 50%;
            background: #dbeafe;
            color: #1d4ed8;
            font-weight: bold;
            font-size: 18px;
        }

        .candidate-name {
            margin: 0 0 5px;
            font-size: 20px;
        }

        .candidate-file {
            color: var(--muted);
            font-size: 13px;
        }

        .score {
            min-width: 120px;
            text-align: center;
            border-radius: 12px;
            padding: 12px 16px;
            background: #ecfdf5;
            color: #166534;
        }

        .score strong {
            display: block;
            font-size: 26px;
        }

        .score span {
            font-size: 12px;
        }

        .progress-wrapper {
            margin: 18px 0;
        }

        .progress-label {
            display: flex;
            justify-content: space-between;
            font-size: 13px;
            color: var(--muted);
            margin-bottom: 6px;
        }

        .progress {
            width: 100%;
            height: 10px;
            border-radius: 20px;
            background: #e2e8f0;
            overflow: hidden;
        }

        .progress-bar {
            height: 100%;
            background: linear-gradient(
                90deg,
                #2563eb,
                #22c55e
            );
            border-radius: inherit;
        }

        .metrics {
            display: grid;
            grid-template-columns: repeat(4, 1fr);
            gap: 10px;
            margin: 18px 0;
        }

        @media (max-width: 800px) {
            .metrics {
                grid-template-columns: repeat(2, 1fr);
            }
        }

        .metric {
            background: #f8fafc;
            border: 1px solid var(--border);
            padding: 12px;
            border-radius: 10px;
        }

        .metric strong {
            display: block;
            font-size: 17px;
            margin-bottom: 4px;
        }

        .metric span {
            color: var(--muted);
            font-size: 12px;
        }

        .details {
            display: grid;
            grid-template-columns: 1fr 1fr;
            gap: 16px;
        }

        @media (max-width: 700px) {
            .details {
                grid-template-columns: 1fr;
            }
        }

        .detail-box {
            border: 1px solid var(--border);
            border-radius: 10px;
            padding: 14px;
        }

        .detail-box h4 {
            margin: 0 0 10px;
        }

        .tag {
            display: inline-block;
            padding: 5px 9px;
            margin: 3px 4px 3px 0;
            background: #eff6ff;
            color: #1d4ed8;
            border-radius: 999px;
            font-size: 12px;
        }

        .tag-danger {
            background: #fef2f2;
            color: #b91c1c;
        }

        .tag-success {
            background: #f0fdf4;
            color: #15803d;
        }

        .notice {
            margin-top: 18px;
            padding: 14px;
            background: #fffbeb;
            border: 1px solid #fde68a;
            border-radius: 10px;
            color: #92400e;
            font-size: 13px;
            line-height: 1.5;
        }

        footer {
            text-align: center;
            color: var(--muted);
            font-size: 13px;
            padding: 20px;
        }

        @media print {
            header,
            form,
            footer,
            .notice {
                display: none !important;
            }

            body {
                background: white;
            }

            .container {
                width: 100%;
                margin: 0;
            }

            .card,
            .candidate-card {
                box-shadow: none;
                break-inside: avoid;
            }
        }
    </style>
</head>

<body>
    <header>
        <div class="header-content">
            <h1>AI Resume Screening System</h1>
            <p>
                Analyze job descriptions and rank resumes using skills,
                education, experience, and NLP-based document similarity.
            </p>
        </div>
    </header>

    <main class="container">
        {% if errors %}
            <div class="alert alert-error">
                <strong>Please fix the following:</strong>
                <ul>
                    {% for error in errors %}
                        <li>{{ error }}</li>
                    {% endfor %}
                </ul>
            </div>
        {% endif %}

        <section class="card">
            <h2>Upload Job and Candidate Resumes</h2>
            <p class="muted">
                Enter the job description and upload one or more resumes.
                Supported resume formats: PDF, DOCX, and TXT.
            </p>

            <form method="POST" enctype="multipart/form-data">
                <div class="grid">
                    <div>
                        <div class="field">
                            <label for="job_description">
                                Job Description
                            </label>

                            <textarea
                                id="job_description"
                                name="job_description"
                                placeholder="Paste the job description here..."
                            >{{ job_text }}</textarea>

                            <div class="hint">
                                Include responsibilities, skills,
                                qualifications, and years of experience.
                            </div>
                        </div>

                        <div class="field">
                            <label for="job_file">
                                Or Upload Job Description
                            </label>

                            <input
                                id="job_file"
                                type="file"
                                name="job_file"
                                accept=".pdf,.docx,.txt"
                            >

                            <div class="hint">
                                The uploaded file is used only when the text
                                field is empty.
                            </div>
                        </div>
                    </div>

                    <div>
                        <div class="field">
                            <label for="resumes">
                                Candidate Resumes
                            </label>

                            <input
                                id="resumes"
                                type="file"
                                name="resumes"
                                accept=".pdf,.docx,.txt"
                                multiple
                                required
                            >

                            <div class="hint">
                                You can select multiple resumes at once.
                            </div>
                        </div>

                        <div class="notice">
                            <strong>Important:</strong>
                            This application is a decision-support tool.
                            Recruiters should review candidates manually and
                            should not make employment decisions based only on
                            the automated score.
                        </div>
                    </div>
                </div>

                <div class="button-row">
                    <button type="submit">
                        Analyze and Rank Candidates
                    </button>

                    {% if results %}
                        <button
                            type="button"
                            class="secondary-button"
                            onclick="window.print()"
                        >
                            Print Results
                        </button>
                    {% endif %}
                </div>
            </form>
        </section>

        {% if processed and results %}
            <section class="card">
                <h2>Screening Summary</h2>
                <p class="muted">
                    Candidates are ranked from highest to lowest relevance.
                </p>

                <div class="summary-grid">
                    <div class="summary-box">
                        <strong>{{ results|length }}</strong>
                        <span>Candidates analyzed</span>
                    </div>

                    <div class="summary-box">
                        <strong>{{ required_skills|length }}</strong>
                        <span>Required skills detected</span>
                    </div>

                    <div class="summary-box">
                        <strong>{{ required_years }}</strong>
                        <span>Required years of experience</span>
                    </div>

                    <div class="summary-box">
                        <strong>{{ required_education }}</strong>
                        <span>Required education</span>
                    </div>
                </div>

                {% if required_skills %}
                    <div style="margin-top: 18px;">
                        <strong>Detected job skills:</strong>
                        <div style="margin-top: 8px;">
                            {% for skill in required_skills %}
                                <span class="tag">{{ skill }}</span>
                            {% endfor %}
                        </div>
                    </div>
                {% endif %}
            </section>

            <section>
                {% for candidate in results %}
                    <article class="candidate-card">
                        <div class="candidate-top">
                            <div class="candidate-info">
                                <div class="rank">
                                    #{{ candidate.rank }}
                                </div>

                                <div>
                                    <h3 class="candidate-name">
                                        {{ candidate.name }}
                                    </h3>

                                    <div class="candidate-file">
                                        {{ candidate.filename }}
                                    </div>
                                </div>
                            </div>

                            <div class="score">
                                <strong>{{ candidate.final_score }}%</strong>
                                <span>Overall Match</span>
                            </div>
                        </div>

                        <div class="progress-wrapper">
                            <div class="progress-label">
                                <span>Candidate relevance</span>
                                <span>
                                    {{ candidate.final_score }}%
                                </span>
                            </div>

                            <div class="progress">
                                <div
                                    class="progress-bar"
                                    style="width: {{ candidate.final_score }}%;"
                                ></div>
                            </div>
                        </div>

                        <div class="metrics">
                            <div class="metric">
                                <strong>
                                    {{ candidate.skill_score }}%
                                </strong>
                                <span>Skill match</span>
                            </div>

                            <div class="metric">
                                <strong>
                                    {{ candidate.semantic_score }}%
                                </strong>
                                <span>NLP similarity</span>
                            </div>

                            <div class="metric">
                                <strong>
                                    {{ candidate.experience_score }}%
                                </strong>
                                <span>Experience match</span>
                            </div>

                            <div class="metric">
                                <strong>
                                    {{ candidate.education_score }}%
                                </strong>
                                <span>Education match</span>
                            </div>
                        </div>

                        <div class="details">
                            <div class="detail-box">
                                <h4>Experience and Education</h4>

                                <p>
                                    <strong>Experience:</strong>
                                    {{ candidate.candidate_years }}
                                    years
                                </p>

                                <p>
                                    <strong>Required:</strong>
                                    {{ candidate.required_years }}
                                    years
                                </p>

                                <p>
                                    <strong>Education:</strong>
                                    {{ candidate.candidate_education }}
                                </p>

                                <p>
                                    <strong>Job requirement:</strong>
                                    {{ candidate.required_education }}
                                </p>
                            </div>

                            <div class="detail-box">
                                <h4>Matched Skills</h4>

                                {% if candidate.matched_skills %}
                                    {% for skill in candidate.matched_skills %}
                                        <span class="tag tag-success">
                                            {{ skill }}
                                        </span>
                                    {% endfor %}
                                {% else %}
                                    <span class="muted">
                                        No known matching skills detected.
                                    </span>
                                {% endif %}
                            </div>

                            <div class="detail-box">
                                <h4>Missing Skills</h4>

                                {% if candidate.missing_skills %}
                                    {% for skill in candidate.missing_skills %}
                                        <span class="tag tag-danger">
                                            {{ skill }}
                                        </span>
                                    {% endfor %}
                                {% else %}
                                    <span class="tag tag-success">
                                        No detected missing skills
                                    </span>
                                {% endif %}
                            </div>

                            <div class="detail-box">
                                <h4>All Detected Candidate Skills</h4>

                                {% if candidate.candidate_skills %}
                                    {% for skill in candidate.candidate_skills %}
                                        <span class="tag">
                                            {{ skill }}
                                        </span>
                                    {% endfor %}
                                {% else %}
                                    <span class="muted">
                                        No known skills detected.
                                    </span>
                                {% endif %}
                            </div>
                        </div>
                    </article>
                {% endfor %}
            </section>
        {% elif processed and not results %}
            <section class="card">
                <h2>No Results</h2>
                <p class="muted">
                    No readable resume content was found. Make sure the files
                    contain selectable text and use PDF, DOCX, or TXT format.
                </p>
            </section>
        {% endif %}
    </main>

    <footer>
        AI Resume Screening System · NLP-based decision-support prototype
    </footer>
</body>
</html>
"""


# -------------------------------------------------------------------
# Routes
# -------------------------------------------------------------------

@app.route("/", methods=["GET", "POST"])
def index():
    errors = []
    results = []
    required_skills = []
    required_years = 0
    required_education = "Not detected"
    job_text = ""
    processed = False

    if request.method == "POST":
        job_text = request.form.get("job_description", "").strip()

        # If no job description text was entered, try the uploaded file.
        if not job_text:
            job_file = request.files.get("job_file")

            if job_file and job_file.filename:
                if not is_allowed_file(job_file.filename):
                    errors.append(
                        "The job description file must be PDF, DOCX, or TXT."
                    )
                else:
                    try:
                        job_text = extract_text_from_file(job_file).strip()
                    except Exception as exc:
                        errors.append(
                            f"Could not read the job description file: {exc}"
                        )

        if not job_text:
            errors.append(
                "Please enter a job description or upload a job description file."
            )

        uploaded_resumes = request.files.getlist("resumes")

        if not uploaded_resumes or not any(
            file.filename for file in uploaded_resumes
        ):
            errors.append("Please upload at least one resume.")

        candidates = []

        for resume_file in uploaded_resumes:
            if not resume_file or not resume_file.filename:
                continue

            filename = secure_filename(resume_file.filename)

            if not is_allowed_file(filename):
                errors.append(
                    f"{filename}: unsupported format. "
                    "Use PDF, DOCX, or TXT."
                )
                continue

            try:
                extracted_text = extract_text_from_file(resume_file).strip()

                if not extracted_text:
                    errors.append(
                        f"{filename}: no readable text was found."
                    )
                    continue

                candidates.append(
                    {
                        "filename": filename,
                        "name": get_candidate_name(
                            filename,
                            extracted_text,
                        ),
                        "text": extracted_text,
                    }
                )

            except Exception as exc:
                errors.append(
                    f"{filename}: could not be processed. {exc}"
                )

        if job_text and candidates:
            analysis = analyze_candidates(job_text, candidates)

            results = analysis["results"]
            required_skills = analysis["required_skills"]
            required_years = analysis["required_years"]
            required_education = analysis["required_education"]

            processed = True

    return render_template_string(
        HTML_TEMPLATE,
        errors=errors,
        results=results,
        required_skills=required_skills,
        required_years=required_years,
        required_education=required_education,
        job_text=job_text,
        processed=processed,
    )


@app.errorhandler(413)
def request_entity_too_large(error):
    return (
        """
        <h2>Uploaded files are too large.</h2>
        <p>The maximum combined upload size is 50 MB.</p>
        <p><a href="/">Return to application</a></p>
        """,
        413,
    )


if __name__ == "__main__":
    port = int(os.environ.get("PORT", 5000))

    app.run(
        host="0.0.0.0",
        port=port,
        debug=True,
    )
